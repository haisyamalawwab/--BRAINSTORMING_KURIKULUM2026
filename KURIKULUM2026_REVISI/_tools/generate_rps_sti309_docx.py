# -*- coding: utf-8 -*-
"""RPS Batch Generator (unit ke-3): MK STI-309 Rekayasa Perangkat Lunak.

Mengisi Template_RPS_GEN_2026_20092026-SWO.docx (12 tabel, 18 tag {{...}} +
25 tag [...]) dengan data resmi Kurikulum SISTEKIN 2026:
  - Identitas MK        : Dok 005 (baris MK no. 9, Sem 3) & Dok 007 Tabel 20.A
  - CPL P4 / KK5        : Dok 003 (rumusan resmi, verbatim)
  - CPMK & Sub-CPMK     : Dok 007 Tabel 20.B & Dok 043 (verbatim, ABCD + Bloom)
  - Rencana 16 pekan    : Dok 007 Tabel 20.C (backbone pekan) x Sub-CPMK Dok 043
  - Skema asesmen       : 4x Titik Baku Dok 008 tipe Teori (T1 20%, UTS 30%,
                          T2 20%, UAS 30%)
  - Boundary Guardrails : Dok 043 (handoff SRS/SDD dari STI-306; cloud ke
                          STI-417; penjaminan mutu ke FST-712)

Keluaran:
  - DOCX/RPS_STI-309_Software_Engineering.docx   (template terisi penuh)
  - RPS_STI-309_Software_Engineering.md          (pendamping markdown)
"""
import copy
import os

from docx import Document

SCRIPT = os.path.dirname(os.path.abspath(__file__))
WORKDIR = os.path.dirname(SCRIPT)
TEMPLATE = os.path.join(WORKDIR, "Template_RPS_GEN_2026_20092026-SWO.docx")
OUT_DOCX = os.path.join(WORKDIR, "DOCX", "RPS_STI-309_Software_Engineering.docx")
OUT_MD = os.path.join(WORKDIR, "RPS_STI-309_Software_Engineering.md")

TGL_SUSUN = "30 September 2026"

# ---------------------------------------------------------------- data inti
SCALARS = {
    "{{namamk}}": "Rekayasa Perangkat Lunak (Software Engineering)",
    "{{kodemk}}": "STI-309",
    "{{rumpunmk}}": "Rekayasa Perangkat Lunak & Platform (Core STI)",
    "{{skst}}": "3",
    "{{sksp}}": "0",
    "{{sms}}": "3 (Ganjil)",
    "{{tglsusun}}": TGL_SUSUN,
    "{{kodeprodi}}": "B4.5.2.RPS-STI309",
    "{{pengembangrps}}": "Tim KBK Rekayasa Perangkat Lunak & Platform",
    "{{koordinatormk}}": "Dr. Roni Wahyu, S.Kom., M.T.",
    "{{ketuaprodi}}": "[Nama Ketua Prodi SISTEKIN]",
    "{{prodi}}": "Sistem dan Teknologi Informasi",
    "{{namadosen}}": "[Nama Dosen Pengampu STI-309]",
    "{{nuptk}}": "[NUPTK Dosen]",
    "{{dosenpengampu}}": "1. [Nama Dosen Pengampu 1] (Koordinator MK)\n2. [Nama Dosen Pengampu 2]",
    # Prasyarat: Dok 005 (identitas, sumber terkuat) = FST-203 Struktur Data dan
    # Algoritma. Dok 007 Tabel 20.A menulis "FST-205 Pemrograman Lanjut" —
    # DISCREPANCY, dilaporkan, tidak diam-diam diperbaiki di dokumen sumber.
    "{{mkprasyarat}}": "FST-203 Struktur Data dan Algoritma",
    "{{deskripsimk}}": (
        "Rekayasa Perangkat Lunak (STI-309) adalah mata kuliah wajib Core STI pada Semester 3 "
        "yang membekali mahasiswa dengan prinsip dan praktik rekayasa perangkat lunak modern. "
        "Mahasiswa memilih dan menjustifikasi model proses pengembangan (Waterfall, V-Model, "
        "Scrum, Kanban) sesuai karakteristik proyek industri; menerapkan arsitektur perangkat "
        "lunak modular (Layered, Client-Server, Event-Driven) dan prinsip SOLID dengan coupling "
        "rendah dan cohesion tinggi; menerapkan design patterns GoF; serta mendeteksi code "
        "smells dan melakukan refactoring. Pembelajaran menekankan kualitas: merancang skenario "
        "pengujian White-Box (basis path, cyclomatic complexity) dan Black-Box (equivalence "
        "partitioning, boundary value analysis), otomasi pengujian (unit testing, TDD, mocking), "
        "evaluasi kualitas perangkat lunak berdasarkan standar ISO/IEC 25010, manajemen "
        "konfigurasi perangkat lunak (Git branching & CI/CD), hingga penyusunan dokumen SQA "
        "Plan dan Test Execution Report. Sesuai Boundary Guardrails Dok 043, mata kuliah ini "
        "menerima dokumen spesifikasi analisis dan desain (SRS & SDD) dari STI-306 Analisis dan "
        "Perancangan Sistem Informasi, serta menyerahkan integrasi arsitektur cloud kepada "
        "STI-417 Komputasi Awan dan penjaminan mutu formal kepada FST-712."
    ),
    "{{bahankajian}}": (
        "1. BK-IS07 Systems Analysis and Design — prinsip rekayasa perangkat lunak, model proses "
        "pengembangan, arsitektur dan desain modular perangkat lunak.\n"
        "2. BK-IT11 Software Development Practices — praktik pengembangan perangkat lunak: clean "
        "code, refactoring, design patterns, pengujian, dan manajemen konfigurasi.\n"
        "Pokok Bahasan: (1) Pengantar RPL, krisis software & SWEBOK; (2) Model Proses "
        "Tradisional: Waterfall, Incremental, V-Model; (3) Model Proses Adaptif: Agile Manifesto, "
        "Scrum Sprint, Kanban; (4) Arsitektur Perangkat Lunak: Monolitik, Layered, Client-Server, "
        "Pipe-Filter; (5) Refactoring & Kepatuhan 5 Prinsip SOLID; (6) Code Smells & Teknik "
        "Refactoring; (7) Design Patterns GoF (Creational, Structural, Behavioral); (8) SQA & "
        "Standar ISO/IEC 25010; (9) White-Box Testing: Flow Graph, Cyclomatic Complexity, "
        "Coverage; (10) Black-Box Testing: Equivalence Partitioning & Boundary Value Analysis; "
        "(11) Otomasi Pengujian: Unit Testing, TDD, Mocking Framework; (12) Integration, "
        "Regression & Stress/Load Testing; (13) Software Configuration Management: Git Branching "
        "& CI/CD; (14) SQA Plan & Test Execution Report."
    ),
}

CPLS = [
    ("P4", "Menguasai prinsip rekayasa perangkat lunak modern (web, mobile, distributed), "
           "manajemen basis data relasional/NoSQL, rekayasa data & visualisasi, serta prinsip "
           "desain interaksi pengalaman pengguna (UI/UX)."),
    ("KK5", "Mampu merancang pengalaman pengguna berbasis riset (UI/UX) serta merekayasa platform "
            "digital modern yang skalabel (Microservices, Web/Mobile, API, SaaS, BPA)."),
]

CPMKS = [
    ("CPMK-1", "Mahasiswa (A) mampu memilih model proses pengembangan perangkat lunak (Waterfall, "
               "V-Model, Scrum, Kanban) (B) yang sesuai dengan karakteristik proyek industri (C) "
               "secara terjustifikasi (D). [C4]"),
    ("CPMK-2", "Mahasiswa (A) mampu menerapkan arsitektur perangkat lunak modular dan prinsip SOLID "
               "(B) pada perancangan subsistem software (C) dengan coupling rendah dan cohesion "
               "tinggi (D). [C3]"),
    ("CPMK-3", "Mahasiswa (A) mampu merancang skenario pengujian perangkat lunak White-Box (Basis "
               "Path, Cyclomatic Complexity) dan Black-Box (BVA, Equivalence Partitioning) (B) "
               "berdasarkan spesifikasi sistem (C) secara komprehensif (D). [C4]"),
    ("CPMK-4", "Mahasiswa (A) mampu mengevaluasi kualitas perangkat lunak (B) berdasarkan standar "
               "ISO/IEC 25010 dan metrik pengujian otomatis (C) secara objektif (D). [C5]"),
]

SUBCPMKS = [
    ("Sub-CPMK-1.1", "Memilih model proses pengembangan perangkat lunak (Waterfall, V-Model, "
                     "Scrum, Kanban) yang tepat. (C4)"),
    ("Sub-CPMK-1.2", "Membandingkan karakteristik model proses pengembangan untuk proyek "
                     "industri. (C4)"),
    ("Sub-CPMK-2.1", "Menerapkan arsitektur perangkat lunak modular dan prinsip SOLID dengan "
                     "coupling rendah. (C3)"),
    ("Sub-CPMK-2.2", "Menerapkan design pattern (GoF) pada perancangan subsistem perangkat "
                     "lunak. (C3)"),
    ("Sub-CPMK-3.1", "Merancang skenario pengujian White-Box (basis path, cyclomatic "
                     "complexity). (C4)"),
    ("Sub-CPMK-3.2", "Merancang skenario pengujian Black-Box (BVA, equivalence partitioning) "
                     "berdasarkan spesifikasi. (C4)"),
    ("Sub-CPMK-4.1", "Mengevaluasi kualitas perangkat lunak berdasarkan standar ISO/IEC "
                     "25010. (C5)"),
    ("Sub-CPMK-4.2", "Mengevaluasi hasil pengujian unit dan integrasi secara otomatis dan "
                     "objektif. (C5)"),
]

KOLERASI = {
    "header": ["CPL", "CPMK-1", "CPMK-2", "CPMK-3", "CPMK-4", "Instrumen & Bobot Asesmen"],
    "rows": [
        ["P4", "\u2713", "\u2713", "", "", "Tugas 1 (20%) + UTS (30%) = 50%"],
        ["KK5", "", "", "\u2713", "\u2713", "Tugas 2 (20%) + UAS (30%) = 50%"],
    ],
    "note": "Skema 4x Titik Evaluasi Baku (Dok 008, tipe Teori): Tugas 1 Pekan 4 (20%), "
            "UTS Pekan 8 (30%), Tugas 2 Pekan 12 (20%), UAS Pekan 16 (30%) — total 100%.",
}

# Pustaka: komposisi baru mengikuti gaya pilot (Dok 007 tidak menyediakan daftar
# pustaka) — wajib validasi KBK/GPM. Penomoran referensi materi: utama 1-5,
# pendukung melanjutkan 6-12.
PUSTAKA_UTAMA = [
    "1. Sommerville, I. (2016). Software Engineering (10th ed.). Boston: Pearson.",
    "2. Pressman, R. S., & Maxim, B. R. (2020). Software Engineering: A Practitioner's Approach "
    "(9th ed.). New York: McGraw-Hill Education.",
    "3. ISO/IEC (2011). ISO/IEC 25010:2011 — Systems and Software Quality Requirements and "
    "Evaluation (SQuaRE): System and Software Quality Models. Geneva: ISO.",
    "4. Bourque, P., & Fairley, R. E. (2014). Guide to the Software Engineering Body of Knowledge "
    "(SWEBOK Guide) Version 3.0. IEEE Computer Society. https://www.computer.org/education/"
    "bodies-of-knowledge/software-engineering",
    "5. Martin, R. C. (2018). Clean Architecture: A Craftsman's Guide to Software Structure and "
    "Design. Boston: Prentice Hall.",
]
PUSTAKA_PENDUKUNG = [
    "6. Martin, R. C. (2008). Clean Code: A Handbook of Agile Software Craftsmanship. Boston: "
    "Prentice Hall.",
    "7. Gamma, E., Helm, R., Johnson, R., & Vlissides, J. (1994). Design Patterns: Elements of "
    "Reusable Object-Oriented Software. Boston: Addison-Wesley.",
    "8. Beck, K. (2003). Test-Driven Development: By Example. Boston: Addison-Wesley.",
    "9. Chacon, S., & Straub, B. (2014). Pro Git (2nd ed.). New York: Apress. https://git-scm.com/book",
    "10. Myers, G. J., Sandler, C., & Badgett, T. (2011). The Art of Software Testing (3rd ed.). "
    "Hoboken: John Wiley & Sons.",
    "11. Fowler, M. (2018). Refactoring: Improving the Design of Existing Code (2nd ed.). Boston: "
    "Addison-Wesley.",
    "12. Schwaber, K., & Sutherland, J. (2020). The Scrum Guide: The Definitive Guide to Scrum. "
    "https://scrumguides.org",
]

# ------------------------------------------------------------- 16 pekan
# (minggu, kemampuan/sub-CPMK, indikator, kriteria&teknik, luring, daring, materi, bobot)
# Backbone topik & metode = Dok 007 Tabel 20.C; kode Sub-CPMK = Dok 043
# (dicocokkan ke pekan); indikator/kriteria diturunkan dari rumusan kemampuan.
WEEKS = [
    (1, "Pengantar MK & Orientasi — Mampu menguraikan krisis software dan standar IEEE SWEBOK (C2)",
        "Peta konsep RPL & SWEBOK tersusun; kontrak belajar dan skema asesmen disepakati",
        "Checklist formatif; observasi partisipasi",
        "Kuliah & Diskusi 150'",
        "Edlink: modul pengantar & kontrak belajar; forum diskusi",
        "Pengantar Rekayasa Perangkat Lunak, Evolusi Kualitas Software, SWEBOK [1][4]", "\u2014"),
    (2, "Sub-CPMK-1.1: Mampu memilih model proses pengembangan yang tepat (C4)",
        "Matriks karakteristik Waterfall, Incremental, V-Model (kelebihan/kekurangan/konteks) benar",
        "Rubrik kompilasi formatif; kuis Edlink",
        "Ceramah & Analisis Kasus 150'",
        "Edlink: studi kasus proyek industri; forum analisis",
        "Model Proses Tradisional: Waterfall, Incremental, V-Model [1][2]", "\u2014"),
    (3, "Sub-CPMK-1.2: Mampu membandingkan karakteristik model proses tradisional dan adaptif "
        "untuk proyek industri (C4)",
        "Pemetaan peran, artefak, dan sprint Scrum/Kanban pada studi kasus tepat",
        "Rubrik kompilasi formatif; presentasi kasus singkat",
        "Case Method Class 150'",
        "Edlink: video Agile/Scrum; tugas analisis kasus",
        "Model Proses Adaptif: Agile Manifesto, Scrum Sprint, Kanban [2][12]", "\u2014"),
    (4, "Sub-CPMK-2.1: Mampu menganalisis pola arsitektur Layered, Client-Server, Event-Driven (C4)",
        "Justifikasi pemilihan pola arsitektur untuk studi kasus diterima (rubrik Tugas 1)",
        "Rubrik analitik Tugas 1",
        "Problem-Solving Class 150'",
        "Edlink: pengumpulan Tugas 1 via Edlink",
        "Arsitektur Perangkat Lunak: Monolitik, Layered, Client-Server, Pipe-Filter [1][5]",
        "Tugas 1: 20% (Problem Solving Model Proses & Arsitektur — evaluasi Sub-CPMK Pekan 2\u20134)"),
    (5, "Sub-CPMK-2.1: Mampu menerapkan prinsip SOLID pada desain kode perangkat lunak (C3)",
        "Kode studi kasus patuh 5 prinsip SOLID tanpa pelanggaran kritis",
        "Rubrik kompilasi formatif; review kode",
        "Case-Based Workshop 150'",
        "Edlink: starter code; commit repo mingguan",
        "Refactoring Kode & Kepatuhan 5 Prinsip Desain SOLID [6][11]", "\u2014"),
    (6, "Sub-CPMK-2.1: Mampu menganalisis code smells dan teknik refactoring sistematis (C4)",
        "Minimal 3 code smell terdeteksi dan direfaktor dengan alasan terdokumentasi",
        "Rubrik kompilasi formatif; demo refactoring",
        "Problem-Based Learning 150'",
        "Edlink: latihan code smell; forum review kode",
        "Mendeteksi Code Smells: Long Method, Large Class, Duplikasi [11]", "\u2014"),
    (7, "Sub-CPMK-2.2: Mampu menerapkan pola desain GoF (Creational, Structural, Behavioral) (C3)",
        "Minimal 2 design pattern GoF diterapkan tepat konteks pada studi kasus",
        "Rubrik kompilasi formatif; presentasi pola",
        "Case Method Class 150'",
        "Edlink: contoh kasus pattern; forum diskusi",
        "Design Patterns dalam Arsitektur Software Skala Besar [7]", "\u2014"),
    "UTS",
    (9, "Sub-CPMK-4.1: Mampu menguraikan 8 karakteristik kualitas standar ISO/IEC 25010 (C2)",
        "8 karakteristik kualitas dipetakan ke contoh produk software nyata",
        "Rubrik kompilasi formatif; kuis Edlink",
        "Kuliah & Diskusi 150'",
        "Edlink: modul SQA; forum diskusi",
        "Software Quality Assurance (SQA) & Standar ISO/IEC 25010 [3]", "\u2014"),
    (10, "Sub-CPMK-3.1: Mampu menghitung Cyclomatic Complexity dan merancang Basis Path Testing (C4)",
        "Flow graph dan V(G) dihitung benar; test case basis path lengkap",
        "Rubrik kompilasi formatif; penilaian latihan",
        "Problem-Solving Class 150'",
        "Edlink: latihan flow graph; kuis",
        "White-Box Testing: Flow Graph, Cyclomatic Complexity, Coverage [10]", "\u2014"),
    (11, "Sub-CPMK-3.2: Mampu merancang uji Black-Box via Equivalence Partitioning dan BVA (C4)",
        "Kelas ekuivalensi dan nilai batas dirancang lengkap dari spesifikasi",
        "Rubrik kompilasi formatif; penilaian latihan",
        "Case-Based Learning 150'",
        "Edlink: spesifikasi kasus uji; forum diskusi",
        "Black-Box Testing: Equivalence Partitioning & Boundary Value Analysis [10]", "\u2014"),
    (12, "Sub-CPMK-3.1 & 4.2: Mampu menerapkan strategi Unit Testing dan Test-Driven Development "
         "(C3)",
        "Siklus TDD (red-green-refactor) didemonstrasikan; unit test lulus",
        "Rubrik analitik Tugas 2",
        "Workshop Koding 150'",
        "Edlink: pengumpulan Tugas 2 via Edlink",
        "Otomasi Pengujian: Unit Testing, TDD, Mocking Framework [8]",
        "Tugas 2: 20% (Studi Kasus Otomasi Pengujian & TDD — evaluasi Sub-CPMK Pekan 9\u201312)"),
    (13, "Sub-CPMK-4.2: Mampu merancang pengujian integrasi dan System Stress Testing (C4)",
        "Skenario integration/regression/stress test dirancang untuk studi kasus",
        "Rubrik kompilasi formatif; review rancangan",
        "Problem-Solving Class 150'",
        "Edlink: studi kasus integrasi; forum",
        "Integration Testing, Regression Testing, Stress/Load Testing [2][10]", "\u2014"),
    (14, "Sub-CPMK-4.2: Mampu menguraikan manajemen konfigurasi perangkat lunak (SCM) dan Git (C2)",
        "Alur branching Git dan pipeline CI/CD konsep dijelaskan benar",
        "Rubrik kompilasi formatif; demonstrasi Git",
        "Ceramah & Demo 150'",
        "Edlink: tutorial Git; demo pipeline",
        "Software Configuration Management: Git Branching & CI/CD Konsep [9]", "\u2014"),
    (15, "Sub-CPMK-4.1 & 4.2: Mampu menyusun Test Plan dan laporan audit kualitas software tim (C5)",
        "Dokumen SQA Plan dan Test Execution Report tersusun dan dipresentasikan",
        "Rubrik kompilasi presentasi; penilaian antar-tim",
        "Presentasi Kelompok 150'",
        "Edlink: unggah dokumen SQA Plan; umpan balik",
        "Penyusunan Dokumen SQA Plan & Test Execution Report [2][3]", "\u2014"),
    "UAS",
]

UTS_CELL = (
    "Indikator: Skor ujian tertulis \u2265 70; jawaban analisis model proses, arsitektur, dan "
    "SOLID memenuhi rubrik.\n"
    "Kriteria & Teknik: Rubrik analitik UTS; ujian tertulis terjadwal (150').\n"
    "Luring: Ujian Tertulis Terjadwal (150').\n"
    "Daring: Edlink — pengumuman ruang ujian & unggah berkas.\n"
    "Materi: Cakupan Mg 1\u20137 — model proses, arsitektur, SOLID, code smells, design patterns "
    "[1][2][5][7].\n"
    "Bobot Penilaian: UTS 30%."
)
UAS_CELL = (
    "Indikator: Skor ujian komprehensif \u2265 70; jawaban desain arsitektur, pengujian, dan audit "
    "kualitas memenuhi rubrik.\n"
    "Kriteria & Teknik: Rubrik analitik UAS; ujian tertulis komprehensif (150').\n"
    "Luring: Ujian Tertulis Akhir (150').\n"
    "Daring: Edlink — pengumuman ruang ujian & unggah berkas.\n"
    "Materi: Cakupan seluruh Sub-CPMK — desain arsitektur, pengujian White/Black Box, SQA "
    "[1][2][3][10].\n"
    "Bobot Penilaian: UAS 30%."
)

# ------------------------------------------------------------- lampiran
def rub(title, header, rows):
    return ("table", title, header, rows)

LAMPIRAN = [
    # [penugasan1] — komposisi baru, perlu validasi KBK/GPM
    [
        ("p", "TUGAS 1 — PROBLEM SOLVING: MODEL PROSES & ARSITEKTUR PERANGKAT LUNAK (Pekan 4, "
              "Bobot 20%)"),
        ("p", "Mahasiswa menganalisis satu studi kasus proyek perangkat lunak industri yang "
              "diberikan dosen (mis. sistem informasi layanan publik), membandingkan model proses "
              "pengembangan, dan merancang arsitektur perangkat lunak modular."),
        ("p", "Spesifikasi wajib: (1) matriks perbandingan minimal 3 model proses (Waterfall, "
              "V-Model, Scrum, Kanban) menurut minimal 4 dimensi (ruang lingkup, keterlibatan "
              "pengguna, toleransi perubahan, risiko); (2) justifikasi pemilihan model proses "
              "terhadap karakteristik proyek; (3) diagram arsitektur (Layered/Client-Server/"
              "Event-Driven) lengkap dengan tanggung jawab tiap lapisan; (4) pemetaan prinsip "
              "SOLID pada rancangan subsistem; (5) laporan maksimal 6 halaman."),
        ("p", "Deliverable: dokumen PDF laporan + diagram. Cakupan evaluasi: Sub-CPMK-1.1, 1.2, "
              "dan 2.1 (Pekan 2\u20134) \u2192 CPMK-1 dan CPMK-2 (CPL P4). Penilaian menggunakan "
              "rubrik analitik berikut. (Instrumen komposisi baru \u2014 perlu validasi KBK/GPM.)"),
    ],
    # [rubriktugas1]
    [
        rub("Rubrik Penilaian Tugas 1 (Total 100)",
            ["Kriteria", "Bobot", "Sangat Baik (86\u2013100)", "Baik (76\u201385)",
             "Cukup (66\u201375)", "Kurang (56\u201365)", "Sangat Kurang (0\u201355)"],
            [["Matriks Perbandingan Model Proses", "20%", "\u2265 3 model \u00d7 \u2265 4 dimensi "
              "lengkap & tepat", "Perbandingan tepat, 1 dimensi kurang", "Perbandingan dasar "
              "berfungsi", "Perbandingan dangkal", "Tidak ada matriks"],
             ["Justifikasi Pemilihan Model Proses", "20%", "Justifikasi terukur & berbasis "
              "karakteristik proyek", "Justifikasi logis, argumen kurang tajam", "Justifikasi "
              "dapat diterima", "Justifikasi lemah", "Tanpa justifikasi"],
             ["Diagram Arsitektur & Demarkasi Lapisan", "25%", "Arsitektur modular tepat, tanggung "
              "jawab lapisan jelas", "Diagram benar, 1 lapisan kabur", "Diagram dasar berfungsi",
              "Lapisan tercampur", "Tidak ada diagram"],
             ["Pemetaan Prinsip SOLID", "20%", "5 prinsip terpetakan tepat ke rancangan",
              "4 prinsip tepat", "3 prinsip tepat", "Pemetaan salah arah", "Tidak memetakan"],
             ["Kualitas Dokumentasi", "15%", "Laporan rapi, sitasi pustaka lengkap",
              "Laporan baik, minor kelalaian", "Laporan dapat dipahami", "Laporan berantakan",
              "Tidak ada laporan"]]),
    ],
    # [soaluts]
    [
        ("p", "UJIAN TENGAH SEMESTER (Pekan 8, Bobot 30%) — UJIAN TERTULIS TERJADWAL (150 MENIT)"),
        ("p", "Cakupan: Mg 1\u20137 — model proses pengembangan, arsitektur perangkat lunak, "
              "prinsip SOLID, code smells & refactoring, design patterns (CPMK-1 s.d. CPMK-2)."),
        ("p", "Soal: (1) analisis kasus: pilih dan justifikasi model proses untuk proyek yang "
              "dideskripsikan (essay); (2) sketsa arsitektur Layered untuk studi kasus beserta "
              "tanggung jawab tiap lapisan; (3) identifikasi pelanggaran SOLID pada potongan kode "
              "yang diberikan dan usulkan refactoring; (4) pilih design pattern GoF yang tepat "
              "untuk dua kasus dan gambarkan strukturnya; (5) daftar code smells pada potongan "
              "kode dan tindakan perbaikannya."),
        ("p", "Ketentuan: ujian tertulis individual, tertutup; kriteria ketuntasan minimal skor "
              "70. (Instrumen komposisi baru \u2014 perlu validasi KBK/GPM.)"),
    ],
    # [rubrikuts]
    [
        rub("Rubrik Penilaian UTS (Total 100)",
            ["Kriteria", "Bobot", "Poin Penuh", "Poin Sebagian", "Poin Minimal", "Tidak Ada"],
            [["Analisis & Pemilihan Model Proses", "25%", "25", "17", "9", "0"],
             ["Diagram Arsitektur Layered", "25%", "25", "17", "9", "0"],
             ["Analisis SOLID & Usulan Refactoring", "20%", "20", "13", "7", "0"],
             ["Penerapan Design Pattern GoF", "15%", "15", "10", "5", "0"],
             ["Deteksi Code Smell", "15%", "15", "10", "5", "0"]]),
    ],
    # [tugas2]
    [
        ("p", "TUGAS 2 — STUDI KASUS OTOMASI PENGUJIAN & TEST-DRIVEN DEVELOPMENT (Pekan 12, "
              "Bobot 20%)"),
        ("p", "Mahasiswa merancang dan mengimplementasikan pengujian otomatis untuk satu modul "
              "studi kasus perangkat lunak menggunakan pendekatan Test-Driven Development, lalu "
              "mengevaluasi hasilnya terhadap standar kualitas."),
        ("p", "Spesifikasi wajib: (1) siklus TDD (red-green-refactor) terdokumentasi untuk "
              "minimal 5 fitur/unit; (2) perhitungan cyclomatic complexity dan perancangan test "
              "case basis path untuk satu fungsi kunci; (3) perancangan kelas ekuivalensi dan "
              "Boundary Value Analysis untuk satu fungsi berbasis spesifikasi; (4) evaluasi "
              "kualitas hasil pengujian terhadap minimal 4 karakteristik ISO/IEC 25010; (5) "
              "versi kode dikelola dengan Git (commit terstruktur)."),
        ("p", "Deliverable: repository Git + laporan pengujian maksimal 6 halaman. Cakupan "
              "evaluasi: Sub-CPMK-3.1, 3.2, 4.1, dan 4.2 (Pekan 9\u201312) \u2192 CPMK-3 dan "
              "CPMK-4 (CPL KK5). Penilaian menggunakan rubrik analitik berikut. (Instrumen "
              "komposisi baru \u2014 perlu validasi KBK/GPM.)"),
    ],
    # [rubriktugas2]
    [
        rub("Rubrik Penilaian Tugas 2 (Total 100)",
            ["Kriteria", "Bobot", "Sangat Baik (86\u2013100)", "Baik (76\u201385)",
             "Cukup (66\u201375)", "Kurang (56\u201365)", "Sangat Kurang (0\u201355)"],
            [["Siklus TDD & Kelengkapan Unit Test", "25%", "Siklus terdokumentasi utuh, "
              "\u2265 5 unit teruji", "Siklus utuh, 4 unit", "3 unit teruji", "Test tidak "
              "berkala", "Tanpa unit test"],
             ["Cyclomatic Complexity & Basis Path", "20%", "V(G) benar, test case lengkap",
              "V(G) benar, test case kurang 1", "V(G) sebagian benar", "Perhitungan salah",
              "Tidak menghitung"],
             ["Kelas Ekuivalensi & BVA", "20%", "Partisi & nilai batas lengkap dan tepat",
              "Partisi tepat, 1\u20132 BVA kurang", "Partisi dasar berfungsi", "Partisi salah",
              "Tidak dirancang"],
             ["Evaluasi Kualitas ISO/IEC 25010", "15%", "\u2265 4 karakteristik dievaluasi "
              "berbukti", "4 karakteristik, bukti kurang", "3 karakteristik", "Evaluasi dangkal",
              "Tidak dievaluasi"],
             ["Kualitas Kode & SCM (Git)", "20%", "Commit terstruktur, kode bersih",
              "Commit teratur, minor isu", "Repo ada dan jalan", "Riwayat commit kacau",
              "Tanpa version control"]]),
    ],
    # [soaluas]
    [
        ("p", "UJIAN AKHIR SEMESTER (Pekan 16, Bobot 30%) — UJIAN TERTULIS KOMPREHENSIF (150 "
              "MENIT)"),
        ("p", "Cakupan: seluruh CPMK — desain arsitektur perangkat lunak, pengujian White-Box dan "
              "Black-Box, serta Software Quality Assurance."),
        ("p", "Soal: (1) studi kasus: rancang arsitektur modular untuk sistem yang dideskripsikan "
              "dan pilih design pattern yang mendukungnya (essay + diagram); (2) soal pengujian "
              "White-Box: bangun flow graph, hitung cyclomatic complexity, dan susun test case "
              "basis path; (3) soal pengujian Black-Box: rancang kelas ekuivalensi dan BVA dari "
              "spesifikasi yang diberikan; (4) audit kualitas: nilai produk software terhadap 8 "
              "karakteristik ISO/IEC 25010 dan usulkan perbaikan; (5) susun kerangka SQA Plan "
              "singkat untuk proyek studi kasus."),
        ("p", "Ketentuan: ujian tertulis individual, tertutup; kriteria ketuntasan minimal skor "
              "70. (Instrumen komposisi baru \u2014 perlu validasi KBK/GPM.)"),
    ],
    # [rubrikuas]
    [
        rub("Rubrik Penilaian UAS (Total 100)",
            ["Kriteria", "Bobot", "Poin"],
            [["Desain arsitektur modular & design pattern", "25%", "25"],
             ["Pengujian White-Box (flow graph, V(G), basis path)", "25%", "25"],
             ["Pengujian Black-Box (equivalence partitioning, BVA)", "20%", "20"],
             ["SQA & audit ISO/IEC 25010", "20%", "20"],
             ["Kerangka SQA Plan & SCM/CI-CD konsep", "10%", "10"]]),
    ],
]

BLOCK_TAGS = ["[penugasan1]", "[rubriktugas1]", "[soaluts]", "[rubrikuts]",
              "[tugas2]", "[rubriktugas2]", "[soaluas]", "[rubrikuas]"]
ALL_BLOCK_TAGS = (["[nocpl]", "[cpl]", "[nocpmk]", "[cpmk]", "[nosubcpmk]", "[subcpmk]",
                   "[tabelkolerasi]", "[pustakautama]", "[pustakapendukung]",
                   "[mingguke]", "[kemampuan_subcpmk]", "[indikator]", "[kriteriateknik]",
                   "[luring]", "[daring]", "[materi]", "[bobot_nilai]"] + BLOCK_TAGS)


# ------------------------------------------------------------------ util
def unique_cells(row):
    seen, out = set(), []
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        out.append(c)
    return out


def set_cell_text(cell, text):
    """Ganti isi sel dengan teks (multibaris), pertahankan format run pertama."""
    paras = cell.paragraphs
    p = paras[0]
    for extra in paras[1:]:
        for r in extra.runs:
            r.text = ""
    template_run = p.runs[0] if p.runs else None
    for r in list(p.runs):
        r.text = ""
    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        if template_run is not None:
            run.font.name = template_run.font.name
            run.font.size = template_run.font.size
            run.font.bold = template_run.font.bold
        if i < len(lines) - 1:
            run.add_break()


def clear_cell(cell):
    for p in cell.paragraphs:
        for r in list(p.runs):
            r.text = ""


def iter_paragraphs(doc):
    def walk_table(t):
        for row in t.rows:
            seen = set()
            for c in row.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                for p in c.paragraphs:
                    yield p
                for nt in c.tables:
                    yield from walk_table(nt)
    yield from doc.paragraphs
    for t in doc.tables:
        yield from walk_table(t)


def replace_in_paragraph(p, mapping):
    """Ganti semua tag {{...}} pada satu paragraf (teks bisa terpecah antar-run)."""
    text = p.text
    hit = False
    for tag, val in mapping.items():
        if tag in text:
            text = text.replace(tag, val)
            hit = True
    if not hit:
        return False
    lines = text.split("\n")
    proto = p.runs[0] if p.runs else None
    for r in list(p.runs):
        r.text = ""
    for i, line in enumerate(lines):
        run = p.add_run(line)
        if proto is not None:
            run.font.name = proto.font.name
            run.font.size = proto.font.size
            run.font.bold = proto.font.bold
        if i < len(lines) - 1:
            run.add_break()
    return True


def scalar_pass(doc):
    n = 0
    for p in iter_paragraphs(doc):
        if replace_in_paragraph(p, SCALARS):
            n += 1
    return n


def add_rubric_table(cell, header, rows):
    tbl = cell.add_table(rows=len(rows) + 1, cols=len(header))
    try:
        tbl.style = "Table Grid"
    except Exception:
        pass
    for j, h in enumerate(header):
        tbl.rows[0].cells[j].text = h
        for r in tbl.rows[0].cells[j].paragraphs[0].runs:
            r.font.bold = True
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = val


def fill_lampiran(doc):
    for ti, content in zip(range(3, 11), LAMPIRAN):
        cell = doc.tables[ti].rows[0].cells[0]
        clear_cell(cell)
        first = True
        for item in content:
            if item[0] == "p":
                if first:
                    p = cell.paragraphs[0]
                    run = p.add_run(item[1])
                    run.font.bold = True
                    first = False
                else:
                    p = cell.add_paragraph()
                    run = p.add_run(item[1])
                    if item[1].startswith(("TUGAS", "UJIAN")):
                        run.font.bold = True
            else:
                _, title, header, rows = item
                p = cell.add_paragraph()
                r = p.add_run(title)
                r.font.bold = True
                add_rubric_table(cell, header, rows)


def clone_row_after(table, src_idx, n):
    anchor = table.rows[src_idx]._tr
    for _ in range(n):
        new_tr = copy.deepcopy(anchor)
        anchor.addnext(new_tr)
        anchor = new_tr


def internal_residual_check(doc):
    """Cek residual sebelum save (T4 langkah 4)."""
    blob = "\n".join(p.text for p in iter_paragraphs(doc))
    return [tag for tag in list(SCALARS) + ALL_BLOCK_TAGS if tag in blob]


def audit_output(path):
    """T5 — verifikasi terprogram: buka ulang keluaran dan audit."""
    doc = Document(path)
    issues = []

    def walk_text(d):
        return "\n".join(p.text for p in iter_paragraphs(d))

    blob = walk_text(doc)

    # 1) zero residual
    residual = [t for t in list(SCALARS) + ALL_BLOCK_TAGS if t in blob]
    print("== 1) ZERO RESIDUAL ==")
    print(f"  18 tag {{{{...}}}} + 25 tag [...] tersisa: {residual if residual else 'NIHIL'}")
    # teks header [ Pustaka ] / [ Estimasi Waktu ] adalah teks header template, bukan tag
    print(f"  teks header '[ Pustaka ]' (bukan tag, dipertahankan): "
          f"{'[ Pustaka ]' in blob}")

    t1, t2 = doc.tables[1], doc.tables[2]

    # 2) audit posisi sel kunci
    print("== 2) AUDIT POSISI ==")
    hdr = [p.text for p in unique_cells(t1.rows[0])[1].paragraphs if p.text.strip()]
    print(f"  T1 R0 header : {hdr}")
    print(f"  T1 R3 identitas: {[c.text for c in unique_cells(t1.rows[3])]}")
    for ri in (7, 8, 12):
        print(f"  T1 R{ri} CPL  : {[c.text[:45] for c in unique_cells(t1.rows[ri])[1:]]}")
    for ri in (14, 17, 18, 19):
        print(f"  T1 R{ri} CPMK : {[c.text[:45] for c in unique_cells(t1.rows[ri])[1:]]}")
    for ri in (21, 28):
        print(f"  T1 R{ri} SubCP: {[c.text[:45] for c in unique_cells(t1.rows[ri])[1:]]}")
    print(f"  T1 R30 kolerasi: {unique_cells(t1.rows[30])[1].text[:90]!r}")
    nested = unique_cells(t1.rows[30])[1].tables
    print(f"  T1 R30 nested tabel: {len(nested)}; baris nested: "
          f"{[ [c.text[:12] for c in r.cells] for r in nested[0].rows ]}")
    print(f"  T1 R31 deskripsi: {unique_cells(t1.rows[31])[1].text[:70]!r}")
    print(f"  T1 R32 bahankajian: {unique_cells(t1.rows[32])[1].text[:60]!r}")
    print(f"  T1 R34 pustaka utama (n baris teks): "
          f"{len(unique_cells(t1.rows[34])[1].text.splitlines())}")
    print(f"  T1 R36 pustaka pendukung (n baris teks): "
          f"{len(unique_cells(t1.rows[36])[1].text.splitlines())}")
    print(f"  T1 R37 dosenpengampu: {unique_cells(t1.rows[37])[1].text[:60]!r}")
    print(f"  T1 R38 mkprasyarat: {unique_cells(t1.rows[38])[1].text!r}")

    week_rows = {3: "Pekan 1", 9: "Pekan 7", 10: "UTS", 11: "Pekan 9", 14: "Pekan 12",
                 15: "Pekan 13", 17: "Pekan 15", 18: "UAS", 19: "Total"}
    for ri, label in week_rows.items():
        cells = unique_cells(t2.rows[ri])
        print(f"  T2 R{ri} ({label}): {[c.text[:38] for c in cells]}")

    print("  Lampiran:")
    for ti in range(3, 11):
        cell = doc.tables[ti].rows[0].cells[0]
        first = cell.paragraphs[0].text[:60]
        print(f"    Tabel {ti + 1}: par0={first!r}; nested rubrik={len(cell.tables)}")

    # 3) konsistensi
    print("== 3) KONSISTENSI ==")
    checks = [
        ("Kode & nama MK (Dok 005/007)", "STI-309" in blob and "Rekayasa Perangkat Lunak" in blob),
        ("SKS teori T= 3", "T= 3" in blob),
        ("SKS praktikum P= 0", "P= 0" in blob),
        ("Semester 3", "3 (Ganjil)" in blob),
        ("Prasyarat = Dok 005 (FST-203 Struktur Data dan Algoritma)",
         "FST-203 Struktur Data dan Algoritma" in blob),
        ("CPL P4 verbatim Dok 003", "Menguasai prinsip rekayasa perangkat lunak modern" in blob),
        ("CPL KK5 verbatim Dok 003", "merekayasa platform digital modern yang skalabel" in blob),
        ("CPMK-1 verbatim Dok 043/007", "memilih model proses pengembangan perangkat lunak" in blob),
        ("CPMK-4 verbatim Dok 043/007", "standar ISO/IEC 25010 dan metrik pengujian otomatis" in blob),
        ("Sub-CPMK-4.2 verbatim Dok 043", "Mengevaluasi hasil pengujian unit dan integrasi" in blob),
        ("Jangkar Tugas 1 Pekan 4 (20%)", "Tugas 1: 20%" in blob),
        ("Jangkar UTS Pekan 8 (30%)", "UTS 30%" in blob),
        ("Jangkar Tugas 2 Pekan 12 (20%)", "Tugas 2: 20%" in blob),
        ("Jangkar UAS Pekan 16 (30%)", "UAS 30%" in blob),
        ("Total bobot 100%", "Total bobot penilaian" in blob and "100%" in blob),
        ("Kolerasi P4 = T1 20% + UTS 30% = 50%", "Tugas 1 (20%) + UTS (30%) = 50%" in blob),
        ("Kolerasi KK5 = T2 20% + UAS 30% = 50%", "Tugas 2 (20%) + UAS (30%) = 50%" in blob),
    ]
    for label, ok in checks:
        print(f"  [{'LULUS' if ok else 'GAGAL'}] {label}")
        if not ok:
            issues.append(label)
    bobot_sum = 20 + 30 + 20 + 30
    print(f"  Neraca skema 4x Teori: 20 + 30 + 20 + 30 = {bobot_sum}% (harus 100)")
    if bobot_sum != 100:
        issues.append("neraca bobot")

    # 4) cek sampah
    print("== 4) CEK SAMPAH ==")
    sampah = ["TEKNIK MESIN", "TEKNIK MESIN".lower(), "Mesin", "{{", "[mingguke]",
              "[kemampuan_subcpmk]", "[penugasan1]", "[rubrikuas]"]
    for s in sampah:
        ada = s in blob
        print(f"  {'TEMUAN' if ada else 'bersih'}: {s!r}")
        if ada and s != "Mesin":
            issues.append(f"sampah: {s}")
    # "Mesin" boleh muncul? — harus tidak; periksa terpisah agar jelas
    if "Mesin" in blob:
        issues.append("sampah: Mesin")

    print("== HASIL AKHIR ==")
    print(f"  Status: {'SEMUA CEK LULUS' if not issues else 'ADA GAGAL: ' + repr(issues)}")
    return 0 if not issues else 1


def write_markdown():
    L = []
    A = L.append
    A("# RPS STI-309 — REKAYASA PERANGKAT LUNAK (*Software Engineering*)")
    A("")
    A("**Program Studi:** S1 Sistem dan Teknologi Informasi (SISTEKIN)  ")
    A("**Fakultas:** Sains dan Teknologi Informasi (FSTI) — Universitas Widyagama Malang  ")
    A(f"**Nomor Dokumen:** 071030.{SCALARS['{{kodeprodi}}']}  ")
    A(f"**Tanggal Penyusunan:** {TGL_SUSUN}  ")
    A("**Sumber:** Template Master RPS Generator FSTI UWG (`Template_RPS_GEN_2026_20092026-SWO.docx`) "
      "— seluruh 18 tag `{{...}}` + 25 tag `[...]` terisi; data dari Dok 003/005/007/008/043.  ")
    A("**Keluaran pendamping:** `DOCX/RPS_STI-309_Software_Engineering.docx`")
    A("")
    A("---")
    A("")
    A("## HALAMAN 1 — IDENTITAS MATA KULIAH")
    A("")
    A("| Atribut | Isi |")
    A("|---|---|")
    A(f"| Mata Kuliah (MK) | {SCALARS['{{namamk}}']} |")
    A(f"| Kode | {SCALARS['{{kodemk}}']} |")
    A(f"| Rumpun MK | {SCALARS['{{rumpunmk}}']} |")
    A(f"| Bobot (sks) | T= {SCALARS['{{skst}}']}, P= {SCALARS['{{sksp}}']} (Total 3 SKS, Tipe Teori) |")
    A(f"| Semester | {SCALARS['{{sms}}']} |")
    A(f"| Tgl Penyusunan | {SCALARS['{{tglsusun}}']} |")
    A(f"| Pengembang RPS | {SCALARS['{{pengembangrps}}']} |")
    A(f"| Koordinator RMK | {SCALARS['{{koordinatormk}}']} |")
    A(f"| Ketua Program Studi | {SCALARS['{{ketuaprodi}}']} |")
    A(f"| Matakuliah Syarat | {SCALARS['{{mkprasyarat}}']} |")
    A(f"| Dosen Pengampu | {SCALARS['{{dosenpengampu}}'].replace(chr(10), '<br>')} |")
    A("")
    A("## HALAMAN 2 — CAPAIAN PEMBELAJARAN (CP)")
    A("")
    A("### CPL-PRODI yang Dibebankan pada MK")
    A("")
    A("| No | Capaian Pembelajaran Lulusan (Dok 003) |")
    A("|:--:|---|")
    for code, text in CPLS:
        A(f"| **{code}** | {text} |")
    A("")
    A("### Capaian Pembelajaran Mata Kuliah (CPMK) — Format ABCD & Bloom (Dok 043/007)")
    A("")
    A("| No | Rumusan CPMK |")
    A("|:--:|---|")
    for code, text in CPMKS:
        A(f"| **{code}** | {text} |")
    A("")
    A("### Tahapan Belajar (Sub-CPMK) — Dok 043")
    A("")
    A("| No | Sub-CPMK |")
    A("|:--:|---|")
    for code, text in SUBCPMKS:
        A(f"| **{code}** | {text} |")
    A("")
    A("### Korelasi CPL terhadap CPMK")
    A("")
    A("| " + " | ".join(KOLERASI["header"]) + " |")
    A("|" + ":--:|" * len(KOLERASI["header"]))
    for row in KOLERASI["rows"]:
        A("| " + " | ".join(v if v else " " for v in row) + " |")
    A("")
    A(KOLERASI["note"])
    A("")
    A("### Deskripsi Singkat MK")
    A("")
    A(SCALARS["{{deskripsimk}}"])
    A("")
    A("### Bahan Kajian: Materi Pembelajaran")
    A("")
    A(SCALARS["{{bahankajian}}"].replace(chr(10), chr(10) + chr(10)))
    A("")
    A("### Pustaka")
    A("")
    A("**Utama:**")
    A("")
    for item in PUSTAKA_UTAMA:
        A(item)
    A("")
    A("**Pendukung:**")
    A("")
    for item in PUSTAKA_PENDUKUNG:
        A(item)
    A("")
    A("---")
    A("")
    A("## HALAMAN 3-4 — RENCANA PEMBELAJARAN 16 PEKAN")
    A("")
    A("| Mg ke- | Kemampuan Akhir Tahapan Belajar (Sub-CPMK) | Indikator | Kriteria & Teknik | "
      "Luring (offline) | Daring (online) | Materi Pembelajaran [Pustaka] | Bobot Penilaian (%) |")
    A("|:--:|---|---|---|---|---|---|:--:|")
    for w in WEEKS:
        if w == "UTS":
            A("| **8** | Evaluasi Tengah Semester (UTS) — Ujian Tertulis Terjadwal: Analisis Model "
              "Proses, Arsitektur Software, & Prinsip SOLID (CPMK-1 s.d. CPMK-2) | "
              + UTS_CELL.replace(chr(10), "<br>") + " |")
        elif w == "UAS":
            A("| **16** | Evaluasi Akhir Semester (UAS) — Ujian Tertulis Komprehensif: Desain "
              "Arsitektur, Pengujian White/Black Box, & SQA (Seluruh CPMK) | "
              + UAS_CELL.replace(chr(10), "<br>") + " |")
        else:
            mg, kem, ind, krit, lur, dar, mat, bobot = w
            A(f"| {mg} | {kem} | {ind} | {krit} | {lur} | {dar} | {mat} | {bobot} |")
    A("| | **Total bobot penilaian** | | | | | | **100%** |")
    A("")
    A("---")
    A("")
    A("## LAMPIRAN — INSTRUMEN ASESMEN")
    A("")
    A("> Seluruh instrumen lampiran (4 soal + 4 rubrik) merupakan **komposisi baru** mengikuti "
      "gaya pilot RPS — Dok 007 tidak menyediakan naskah soal/rubrik; wajib validasi KBK/GPM.")
    A("")
    for content in LAMPIRAN:
        for item in content:
            if item[0] == "p":
                A(item[1])
                A("")
            else:
                _, title, header, rows = item
                A(f"**{title}**")
                A("")
                A("| " + " | ".join(header) + " |")
                A("|" + "---|" * len(header))
                for row in rows:
                    A("| " + " | ".join(row) + " |")
                A("")
    A("---")
    A("")
    A("## PENGESAHAN")
    A("")
    A("| Memvalidasi, Unit Penjaminan Mutu FSTI | Malang, " + TGL_SUSUN + " — Dosen Pengampu |")
    A("|---|---|")
    A(f"| Nama Pejabat UPM: [Nama Kepala UPM FSTI] | {SCALARS['{{namadosen}}']} / NUPTK. "
      f"{SCALARS['{{nuptk}}']} |")
    A("")
    A(f"Mengesahkan, Ketua Program Studi {SCALARS['{{prodi}}']}: {SCALARS['{{ketuaprodi}}']}")
    A("")
    A("---")
    A("")
    A("*RPS STI-309 dibangkitkan dari Template Master RPS Generator FSTI UWG oleh "
      "`_tools/generate_rps_sti309_docx.py` — unit ke-3 RPS Batch Generator "
      "(setelah STI-416 V2 & STI-311).*")

    with open(OUT_MD, "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    print(f"Keluaran Markdown    : {OUT_MD}")


def main():
    doc = Document(TEMPLATE)
    t1, t2 = doc.tables[1], doc.tables[2]

    # 0) perbaikan header institusi (warisan template sumber, keluaran saja)
    header_cell = unique_cells(t1.rows[0])[1]
    hdr_paras = [p for p in header_cell.paragraphs if p.text.strip()]
    for r in list(hdr_paras[1].runs):
        r.text = ""
    hdr_paras[1].add_run("FAKULTAS SAINS DAN TEKNOLOGI INFORMASI")
    for r in list(hdr_paras[2].runs):
        r.text = ""
    hdr_paras[2].add_run("PROGRAM STUDI S1 SISTEM DAN TEKNOLOGI INFORMASI")

    # 1) CPL: baris 7-12 -> isi 2, kosongkan 4
    for i in range(6):
        cells = unique_cells(t1.rows[i + 7])
        if i < len(CPLS):
            set_cell_text(cells[1], CPLS[i][0])
            set_cell_text(cells[2], CPLS[i][1])
        else:
            clear_cell(cells[1])
            clear_cell(cells[2])

    # 2) CPMK: baris 14-19 -> isi 4, kosongkan 2
    for i in range(6):
        cells = unique_cells(t1.rows[i + 14])
        if i < len(CPMKS):
            set_cell_text(cells[1], CPMKS[i][0])
            set_cell_text(cells[2], CPMKS[i][1])
        else:
            clear_cell(cells[1])
            clear_cell(cells[2])

    # 3) Sub-CPMK: baris 21-25 (5) + klon 3 -> isi 8
    clone_row_after(t1, 25, 3)
    assert "[tabelkolerasi]" in unique_cells(t1.rows[30])[1].text, "indeks kolerasi bergeser!"
    for i in range(8):
        cells = unique_cells(t1.rows[i + 21])
        set_cell_text(cells[1], SUBCPMKS[i][0])
        set_cell_text(cells[2], SUBCPMKS[i][1])

    # 4) Tabel kolerasi (baris 30 = 27 asli + 3 baris klon Sub-CPMK)
    kol_cell = unique_cells(t1.rows[30])[1]
    assert "[tabelkolerasi]" in kol_cell.text, "indeks baris kolerasi bergeser!"
    clear_cell(kol_cell)
    kol_cell.paragraphs[0].add_run("Matriks Korelasi CPL terhadap CPMK dan Bobot Asesmen:")
    add_rubric_table(kol_cell, KOLERASI["header"], KOLERASI["rows"])
    kol_cell.add_paragraph(KOLERASI["note"])

    # 5) pustaka (indeks pasca-klon: 31=deskripsi, 32=bahankajian, 34=utama, 36=pendukung)
    assert "{{deskripsimk}}" in unique_cells(t1.rows[31])[1].text, "indeks deskripsi bergeser!"
    assert "{{bahankajian}}" in unique_cells(t1.rows[32])[1].text, "indeks bahan kajian bergeser!"
    set_cell_text(unique_cells(t1.rows[34])[1], "\n".join(PUSTAKA_UTAMA))
    set_cell_text(unique_cells(t1.rows[36])[1], "\n".join(PUSTAKA_PENDUKUNG))

    # 6) Tabel 16 pekan: baris 3-9 = pekan 1-7; baris 10 = UTS; baris 11-14 = pekan 9-12;
    #    klon 3 baris untuk pekan 13-15; baris UAS bergeser ke 18.
    clone_row_after(t2, 14, 3)
    weeks = [w for w in WEEKS if not isinstance(w, str)]
    tag_rows = list(range(3, 10)) + list(range(11, 15)) + [15, 16, 17]  # 14 baris ter-tag
    assert len(tag_rows) == len(weeks) == 14
    for row_idx, week in zip(tag_rows, weeks):
        cells = unique_cells(t2.rows[row_idx])
        for cell, val in zip(cells, week):
            set_cell_text(cell, str(val))
    # UTS (baris 10) dan UAS (baris 18 pasca-klon)
    uts_cells = unique_cells(t2.rows[10])
    set_cell_text(uts_cells[1], "Evaluasi Tengah Semester (UTS) — Ujian Tertulis Terjadwal: "
                                "Analisis Model Proses, Arsitektur Software, & Prinsip SOLID "
                                "(CPMK-1 s.d. CPMK-2)")
    set_cell_text(uts_cells[2], UTS_CELL)
    uas_cells = unique_cells(t2.rows[18])
    set_cell_text(uas_cells[1], "Evaluasi Akhir Semester (UAS) — Ujian Tertulis Komprehensif: "
                                "Desain Arsitektur, Pengujian White/Black Box, & SQA "
                                "(Seluruh CPMK)")
    set_cell_text(uas_cells[2], UAS_CELL)

    # 7) Lampiran (tabel 3-10)
    fill_lampiran(doc)

    # 8) tanggal undangan tanda tangan & cover
    for p in iter_paragraphs(doc):
        if p.text.strip() == "Malang, Tanggal Bulan Tahun":
            for r in list(p.runs):
                r.text = ""
            p.add_run(f"Malang, {TGL_SUSUN}")
        if p.text.strip() == "AGUSTUS, 2026":
            for r in list(p.runs):
                r.text = ""
            p.add_run("SEPTEMBER, 2026")

    # 9) scalar pass terakhir (semua {{tag}} termasuk yang ada di tabel nested)
    n_scalar = scalar_pass(doc)

    # 10) cek residual internal sebelum save
    residual = internal_residual_check(doc)
    if residual:
        raise SystemExit(f"RESIDUAL SEBELUM SAVE: {residual}")

    os.makedirs(os.path.dirname(OUT_DOCX), exist_ok=True)
    doc.save(OUT_DOCX)

    print(f"Scalar terisi        : {n_scalar} paragraf")
    print(f"Keluaran DOCX        : {OUT_DOCX}")
    write_markdown()

    return audit_output(OUT_DOCX)


if __name__ == "__main__":
    raise SystemExit(main())
